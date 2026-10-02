"""Extracción local de texto de adjuntos (camino B).

El texto se extrae en el propio servicio (``pypdf`` para PDF, ``python-docx``
para Word) y se concatena al prompt con un separador explícito. Así el flujo no
depende de la Files API de ningún proveedor — el fallback a Gemini sigue
funcionando igual — y deja el texto listo para el chunking de RAG del módulo 3.
El precio: se pierden imágenes y diagramas, y un PDF escaneado no tiene texto.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
EMPTY_MARKER = "[sin texto extraíble]"


class AttachmentError(Exception):
    """Error de adjunto con el código HTTP que debe devolver el endpoint."""

    def __init__(self, status_code: int, detail: str) -> None:
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


@dataclass
class ExtractedAttachment:
    filename: str
    text: str
    chars: int  # caracteres extraídos antes de truncar
    truncated: bool = False
    empty: bool = False

    @property
    def warnings(self) -> list[str]:
        if self.empty:
            return [f"'{self.filename}' no contiene texto extraíble (¿PDF escaneado?)."]
        if self.truncated:
            return [f"'{self.filename}' se ha truncado: tenía {self.chars} caracteres."]
        return []


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    # Muchas especificaciones ponen los requisitos en tablas.
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def extract_attachment(
    filename: str, data: bytes, max_bytes: int, max_chars: int
) -> ExtractedAttachment:
    """Extrae y acota el texto de un adjunto.

    Raises:
        AttachmentError: 415 si la extensión no es .pdf/.docx, 413 si supera
            ``max_bytes``, 422 si el archivo no se puede leer.
    """
    filename = filename or "sin_nombre"
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise AttachmentError(
            415, f"Formato no soportado: '{filename}'. Solo se aceptan PDF (.pdf) y Word (.docx)."
        )
    if len(data) > max_bytes:
        raise AttachmentError(
            413,
            f"'{filename}' ocupa {len(data) / 1_000_000:.1f} MB; "
            f"el máximo es {max_bytes / 1_000_000:.1f} MB.",
        )

    try:
        text = _pdf_text(data) if extension == ".pdf" else _docx_text(data)
    except Exception as exc:  # noqa: BLE001 — pypdf/docx lanzan tipos muy variados
        raise AttachmentError(422, f"No se pudo leer '{filename}': {exc}") from exc

    text = text.strip()
    if not text:
        return ExtractedAttachment(filename=filename, text=EMPTY_MARKER, chars=0, empty=True)

    chars = len(text)
    if chars > max_chars:
        omitted = chars - max_chars
        text = f"{text[:max_chars]}\n[…truncado: se omitieron {omitted} caracteres]"
        return ExtractedAttachment(filename=filename, text=text, chars=chars, truncated=True)

    return ExtractedAttachment(filename=filename, text=text, chars=chars)


def format_attachments(attachments: list[ExtractedAttachment]) -> str:
    """Bloques ``--- attachment: nombre ---`` listos para el prompt."""
    return "\n\n".join(f"--- attachment: {a.filename} ---\n{a.text}" for a in attachments)
