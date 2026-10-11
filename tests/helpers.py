"""Utilidades compartidas por los tests: documentos generados en memoria y LLM falso."""

from __future__ import annotations

import io
import json


def make_pdf(text: str) -> bytes:
    """PDF mínimo de una página con ``text`` (ASCII, sin paréntesis).

    Con ``text=""`` genera una página sin texto: simula un PDF escaneado.
    Se construye a mano para no añadir dependencias solo para tests.
    """
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1") if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


KNOWN_TECHNOLOGIES = ["React", "FastAPI", "PostgreSQL", "Kubernetes"]


def fake_result_json(summary: str, total_min: int = 100, total_max: int = 150) -> str:
    """JSON de EstimationResult válido con una sola fase (o descuadrado si se pide)."""
    return json.dumps(
        {
            "summary": summary,
            "confidence_pct": 70,
            "phases": [
                {"name": "Desarrollo", "hours_min": 100, "hours_max": 150,
                 "confidence_pct": 70, "summary": "Desarrollo del alcance descrito."}
            ],
            "total_hours_min": total_min,
            "total_hours_max": total_max,
            "total_duration_weeks": 4,
            "assumptions": [],
            "out_of_scope": [],
        }
    )


ACCEPT_JSON = json.dumps({"verdict": "aceptar", "issues": [], "confidence_in_review": 80})


def critic_json(verdict: str, *severities: str) -> str:
    """JSON de CriticFeedback con una issue por gravedad dada."""
    issues = [
        {"category": "fase_desproporcionada", "severity": s, "field_path": "phases[0].hours_max",
         "description": "La fase de desarrollo concentra todo el esfuerzo.",
         "suggested_fix": "Separa backend y frontend."}
        for s in severities
    ]
    return json.dumps({"verdict": verdict, "issues": issues, "confidence_in_review": 70})


class FakeLLM:
    """Sustituto de ``LLMWrapper``: graba cada llamada y responde sin red.

    - Extractor (system con "extractor de hechos"): JSON de hechos por palabras
      clave del mensaje del cliente, o ``extractor_text`` literal.
    - Resumidor (system con "resumidor"): un resumen corto que NO copia el texto
      de los turnos, o error si ``fail_summarizer``.
    - Moderador (system con "moderador de contenido"): ``{"flagged": ...}``, o
      error si ``fail_moderator``.
    - Crítico (system con "auditor independiente"): ``critic_responses`` en cola o
      "aceptar"; error si ``fail_critic``.
    - Estimador: ``structured_responses`` en cola si se dan; si no, una
      estimación cuyo summary nombra las tecnologías del mensaje actual (así un
      adjunto influye en la salida), como JSON si se pidió ``json_schema``.
    """

    def __init__(
        self,
        extractor_text: str | None = None,
        fail_estimator: bool = False,
        estimator_text: str | None = None,
        structured_responses: list[str] | None = None,
        fail_summarizer: bool = False,
        moderation_flagged: bool = False,
        fail_moderator: bool = False,
        moderator_text: str | None = None,
        critic_responses: list[str] | None = None,
        fail_critic: bool = False,
    ):
        self.calls: list[list[dict]] = []
        self.kinds: list[str] = []
        self.extractor_text = extractor_text
        self.fail_estimator = fail_estimator
        self.estimator_text = estimator_text
        self.structured_responses = list(structured_responses or [])
        self.fail_summarizer = fail_summarizer
        self.moderation_flagged = moderation_flagged
        self.fail_moderator = fail_moderator
        self.moderator_text = moderator_text
        self.critic_responses = list(critic_responses or [])
        self.fail_critic = fail_critic

    @staticmethod
    def _kind(messages: list[dict]) -> str:
        system = messages[0]["content"]
        if "moderador de contenido" in system:
            return "moderator"
        if "auditor independiente" in system:
            return "critic"
        if "extractor de hechos" in system:
            return "extractor"
        if "resumidor" in system:
            return "summarizer"
        return "estimator"

    @property
    def estimator_calls(self) -> list[list[dict]]:
        return [c for c, k in zip(self.calls, self.kinds) if k == "estimator"]

    @property
    def critic_calls(self) -> list[list[dict]]:
        return [c for c, k in zip(self.calls, self.kinds) if k == "critic"]

    @property
    def summarizer_calls(self) -> list[list[dict]]:
        return [c for c, k in zip(self.calls, self.kinds) if k == "summarizer"]

    def complete(self, messages: list[dict], max_tokens: int = 4000, json_schema=None) -> dict:
        kind = self._kind(messages)
        self.calls.append(messages)
        self.kinds.append(kind)
        last = messages[-1]["content"]
        base = {"model": "fake-model", "provider": "fake", "fallback_used": False}

        if kind == "moderator":
            if self.fail_moderator:
                raise RuntimeError("moderador caído")
            if self.moderator_text is not None:
                return {**base, "text": self.moderator_text}
            categories = ["violencia"] if self.moderation_flagged else []
            return {**base, "text": json.dumps(
                {"flagged": self.moderation_flagged, "categories": categories})}

        if kind == "critic":
            if self.fail_critic:
                raise RuntimeError("crítico caído")
            if self.critic_responses:
                return {**base, "text": self.critic_responses.pop(0)}
            return {**base, "text": ACCEPT_JSON}

        if kind == "extractor":
            if self.extractor_text is not None:
                return {**base, "text": self.extractor_text}
            # Solo mira el mensaje del cliente y los adjuntos, no los hechos previos.
            relevant = last.split("<respuesta_estimador>")[0].split("</hechos_previos>")[-1]
            facts = {
                "project_name": "Hotelia" if "Hotelia" in relevant else None,
                "assumed_team_size": 4 if "seremos 4" in relevant.lower() else None,
                "mentioned_technologies": [t for t in KNOWN_TECHNOLOGIES if t in relevant],
                "agreed_scope": None,
            }
            return {**base, "text": json.dumps(facts)}

        if kind == "summarizer":
            if self.fail_summarizer:
                raise RuntimeError("resumidor caído")
            n = len(self.summarizer_calls)
            return {**base, "text": f"Resumen acumulado número {n} de la conversación."}

        if self.fail_estimator:
            raise RuntimeError("proveedor caído")
        if self.structured_responses:
            return {**base, "text": self.structured_responses.pop(0)}
        if self.estimator_text is not None:
            return {**base, "text": self.estimator_text}
        # El turno actual es el último mensaje del cliente que no sea una
        # petición de corrección del bucle de reintento.
        user_texts = [m["content"] for m in messages
                      if m["role"] == "user"
                      and not m["content"].startswith(("Tu JSON no cumple", "Un revisor independiente"))]
        techs = [t for t in KNOWN_TECHNOLOGIES if t in user_texts[-1]] or ["sin stack definido"]
        n = len(self.estimator_calls)
        summary = f"Estimación #{n}. Stack: {', '.join(techs)}."
        if json_schema is None:
            # Hasta Task 10, run_turn llama sin esquema y espera texto libre.
            return {**base, "text": summary}
        return {**base, "text": fake_result_json(summary)}
